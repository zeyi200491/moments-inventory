from __future__ import annotations

import json
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

import openpyxl


class ModelCatalogError(ValueError):
    pass


@dataclass(slots=True)
class ModelCatalogAlias:
    alias: str
    normalized_key: str
    brand: str
    brand_en: str
    canonical_model: str
    safe: bool
    ambiguous: bool
    source_row: int


@dataclass(slots=True)
class ModelCatalogEntry:
    brand: str
    brand_en: str
    canonical_model: str
    safe_aliases: list[str]
    all_aliases: list[str]
    ambiguous_aliases: list[str]
    safe_standard_key: str
    source_row: int
    safe_standard_keys: list[str] = field(default_factory=list)


@dataclass(slots=True)
class ModelCatalogIndex:
    version: int
    models: list[ModelCatalogEntry]
    aliases_by_key: dict[str, list[ModelCatalogAlias]]


REQUIRED_COLUMNS = [
    "品牌",
    "品牌英文名",
    "规范型号/物理机型",
    "有歧义名称",
    "安全搜索匹配名称合集",
    "全部搜索匹配名称合集",
    "安全标准化匹配键",
]


_MODEL_KEY_PATTERN = re.compile(r"[^0-9a-zA-Z\u4e00-\u9fffμ]+")


def normalize_model_key(value: str) -> str:
    return _MODEL_KEY_PATTERN.sub("", str(value).strip().lower())


def build_model_catalog_index(excel_path: str | Path) -> ModelCatalogIndex:
    workbook = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    try:
        if "完整型号主表" not in workbook.sheetnames:
            raise ModelCatalogError("缺少工作表: 完整型号主表")

        sheet = workbook["完整型号主表"]
        rows = sheet.iter_rows(values_only=True)
        header_row = next(rows, None)
        header_map = _build_header_map(header_row)

        models: list[ModelCatalogEntry] = []
        aliases_by_key: dict[str, list[ModelCatalogAlias]] = {}

        for row_number, row in enumerate(rows, start=2):
            brand = _cell_text(_row_value(row, header_map, "品牌"))
            canonical_model = _cell_text(_row_value(row, header_map, "规范型号/物理机型"))
            if not brand or not canonical_model:
                continue

            brand_en = _cell_text(_row_value(row, header_map, "品牌英文名"))
            safe_aliases = _split_aliases(
                _cell_text(_row_value(row, header_map, "安全搜索匹配名称合集"))
            )
            all_aliases = _split_aliases(
                _cell_text(_row_value(row, header_map, "全部搜索匹配名称合集"))
            )
            searchable_aliases = _merge_aliases(safe_aliases, all_aliases)
            ambiguous_aliases = _split_aliases(
                _cell_text(_row_value(row, header_map, "有歧义名称"))
            )
            safe_standard_keys = _normalize_alias_keys(
                _split_aliases(
                    _cell_text(_row_value(row, header_map, "安全标准化匹配键"))
                )
            )
            safe_standard_key = safe_standard_keys[0] if safe_standard_keys else ""

            entry = ModelCatalogEntry(
                brand=brand,
                brand_en=brand_en,
                canonical_model=canonical_model,
                safe_aliases=safe_aliases,
                all_aliases=all_aliases,
                ambiguous_aliases=ambiguous_aliases,
                safe_standard_key=safe_standard_key,
                source_row=row_number,
                safe_standard_keys=safe_standard_keys,
            )
            models.append(entry)

            safe_keys = {normalize_model_key(alias) for alias in safe_aliases}
            ambiguous_keys = {normalize_model_key(alias) for alias in ambiguous_aliases}
            for alias in searchable_aliases:
                normalized_key = normalize_model_key(alias)
                if not normalized_key:
                    continue
                aliases_by_key.setdefault(normalized_key, []).append(
                    ModelCatalogAlias(
                        alias=alias,
                        normalized_key=normalized_key,
                        brand=brand,
                        brand_en=brand_en,
                        canonical_model=canonical_model,
                        safe=normalized_key in safe_keys,
                        ambiguous=normalized_key in ambiguous_keys,
                        source_row=row_number,
                    )
                )

        return ModelCatalogIndex(version=1, models=models, aliases_by_key=aliases_by_key)
    finally:
        workbook.close()


def save_model_catalog_index(index: ModelCatalogIndex, output_path: str | Path) -> None:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(asdict(index), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def load_model_catalog_index(path: str | Path) -> ModelCatalogIndex:
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    return ModelCatalogIndex(
        version=data.get("version", 1),
        models=[_model_entry_from_dict(model) for model in data.get("models", [])],
        aliases_by_key={
            key: [_model_alias_from_dict(alias) for alias in aliases]
            for key, aliases in data.get("aliases_by_key", {}).items()
        },
    )


def _build_header_map(header_row: tuple[Any, ...] | None) -> dict[str, int]:
    headers = [_cell_text(cell) for cell in (header_row or ())]
    header_map = {header: index for index, header in enumerate(headers) if header}
    missing_columns = [column for column in REQUIRED_COLUMNS if column not in header_map]
    if missing_columns:
        raise ModelCatalogError(f"缺少必需列: {', '.join(missing_columns)}")
    return header_map


def _row_value(row: tuple[Any, ...], header_map: dict[str, int], header: str) -> Any:
    index = header_map[header]
    if index >= len(row):
        return None
    return row[index]


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _split_aliases(value: str) -> list[str]:
    aliases: list[str] = []
    seen: set[str] = set()
    for alias in re.split(r"[;；]", value):
        alias = alias.strip()
        if alias and alias not in seen:
            aliases.append(alias)
            seen.add(alias)
    return aliases


def _merge_aliases(*alias_groups: list[str]) -> list[str]:
    aliases: list[str] = []
    seen: set[str] = set()
    for alias_group in alias_groups:
        for alias in alias_group:
            if not normalize_model_key(alias) or alias in seen:
                continue
            aliases.append(alias)
            seen.add(alias)
    return aliases


def _normalize_alias_keys(aliases: list[str]) -> list[str]:
    keys: list[str] = []
    seen: set[str] = set()
    for alias in aliases:
        key = normalize_model_key(alias)
        if key and key not in seen:
            keys.append(key)
            seen.add(key)
    return keys


def _model_entry_from_dict(value: dict[str, Any]) -> ModelCatalogEntry:
    if "safe_standard_keys" in value:
        safe_standard_keys = _normalize_alias_keys(
            [str(key) for key in value.get("safe_standard_keys", [])]
        )
    elif value.get("safe_standard_key"):
        safe_standard_keys = _normalize_alias_keys([str(value.get("safe_standard_key", ""))])
    else:
        safe_standard_keys = []
    safe_standard_key = safe_standard_keys[0] if safe_standard_keys else ""

    return ModelCatalogEntry(
        brand=str(value.get("brand", "")),
        brand_en=str(value.get("brand_en", "")),
        canonical_model=str(value.get("canonical_model", "")),
        safe_aliases=list(value.get("safe_aliases", [])),
        all_aliases=list(value.get("all_aliases", [])),
        ambiguous_aliases=list(value.get("ambiguous_aliases", [])),
        safe_standard_key=safe_standard_key,
        source_row=int(value.get("source_row", 0) or 0),
        safe_standard_keys=safe_standard_keys,
    )


def _model_alias_from_dict(value: dict[str, Any]) -> ModelCatalogAlias:
    return ModelCatalogAlias(
        alias=str(value.get("alias", "")),
        normalized_key=str(value.get("normalized_key", "")),
        brand=str(value.get("brand", "")),
        brand_en=str(value.get("brand_en", "")),
        canonical_model=str(value.get("canonical_model", "")),
        safe=bool(value.get("safe", False)),
        ambiguous=bool(value.get("ambiguous", False)),
        source_row=int(value.get("source_row", 0) or 0),
    )
