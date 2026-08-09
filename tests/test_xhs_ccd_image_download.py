from pathlib import Path
import importlib.util
import sys

import openpyxl
import pytest


SCRIPT_PATH = Path(__file__).resolve().parents[1] / "scripts" / "download_xhs_ccd_images.py"
spec = importlib.util.spec_from_file_location("download_xhs_ccd_images", SCRIPT_PATH)
xhs = importlib.util.module_from_spec(spec)
assert spec.loader is not None
sys.modules[spec.name] = xhs
spec.loader.exec_module(xhs)


def make_workbook(path: Path, rows: list[list[str]]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CCD完整型号表"
    ws.append(["序号", "品牌", "统一物理机型/规范名", "产品类别", "系列", "欧亚/中国常见名", "传感器"])
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


def test_parse_args_accepts_excel_path(tmp_path: Path):
    excel_path = tmp_path / "models.xlsx"
    excel_path.write_bytes(b"placeholder")

    args = xhs.parse_args([str(excel_path)])

    assert args.excel_path == excel_path
    assert args.output_dir == xhs.OUTPUT_ROOT
    assert args.profile_dir == xhs.DEFAULT_PROFILE_DIR


def test_parse_args_rejects_missing_file(tmp_path: Path):
    missing_path = tmp_path / "missing.xlsx"

    with pytest.raises(SystemExit) as exc:
        xhs.parse_args([str(missing_path)])

    assert exc.value.code == 2


def test_parse_args_rejects_non_xlsx_file(tmp_path: Path, capsys):
    txt_path = tmp_path / "models.txt"
    txt_path.write_text("not an excel file", encoding="utf-8")

    with pytest.raises(SystemExit) as exc:
        xhs.parse_args([str(txt_path)])

    assert exc.value.code == 2
    assert "xlsx" in capsys.readouterr().err.lower()


def test_parse_args_strips_matching_outer_quotes_from_excel_path(tmp_path: Path):
    xlsx_path = tmp_path / "models.xlsx"
    xlsx_path.write_bytes(b"placeholder")

    double_quoted_args = xhs.parse_args([f'"{xlsx_path}"'])
    single_quoted_args = xhs.parse_args([f"'{xlsx_path}'"])

    assert double_quoted_args.excel_path == xlsx_path
    assert single_quoted_args.excel_path == xlsx_path


def test_load_model_tasks_prefers_common_name_and_filters_ccd(tmp_path: Path):
    workbook = make_workbook(
        tmp_path / "canon.xlsx",
        [
            ["1", "Canon 佳能", "PowerShot A720 IS", "相机", "A", "佳能A720", "CCD"],
            ["2", "Canon 佳能", "PowerShot G7 X", "相机", "G", "佳能G7X", "CMOS"],
            ["3", "Canon 佳能", "PowerShot A570 IS", "相机", "A", "", "ccd"],
        ],
    )

    tasks = xhs.load_model_tasks(workbook)

    assert [task.model_name for task in tasks] == ["佳能A720", "PowerShot A570 IS"]
    assert [task.search_query for task in tasks] == ["佳能A720 出片图", "PowerShot A570 IS 出片图"]
    assert {task.brand_dir for task in tasks} == {"佳能"}


def test_load_model_tasks_deduplicates_by_brand_and_model(tmp_path: Path):
    workbook = make_workbook(
        tmp_path / "canon.xlsx",
        [
            ["1", "Canon 佳能", "PowerShot A720 IS", "相机", "A", "佳能A720", "CCD"],
            ["2", "Canon 佳能", "PowerShot A720 IS", "相机", "A", "佳能A720", "CCD"],
        ],
    )

    tasks = xhs.load_model_tasks(workbook)

    assert len(tasks) == 1
    assert tasks[0].model_name == "佳能A720"


def test_load_model_tasks_reports_missing_headers(tmp_path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["品牌", "传感器"])
    path = tmp_path / "bad.xlsx"
    wb.save(path)

    with pytest.raises(xhs.ExcelFormatError) as exc:
        xhs.load_model_tasks(path)

    assert "统一物理机型/规范名" in str(exc.value)
    assert "欧亚/中国常见名" in str(exc.value)


def test_load_model_tasks_closes_workbook_after_success(tmp_path: Path):
    workbook = make_workbook(
        tmp_path / "canon.xlsx",
        [
            ["1", "Canon 佳能", "PowerShot A720 IS", "相机", "A", "佳能A720", "CCD"],
        ],
    )

    xhs.load_model_tasks(workbook)

    workbook.unlink()
    assert not workbook.exists()


def test_load_model_tasks_reports_duplicate_headers(tmp_path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["品牌", "品牌", "统一物理机型/规范名", "欧亚/中国常见名", "传感器"])
    path = tmp_path / "duplicate.xlsx"
    wb.save(path)

    with pytest.raises(xhs.ExcelFormatError) as exc:
        xhs.load_model_tasks(path)

    assert "Excel 存在重复字段：品牌" in str(exc.value)


def test_safe_path_part_prefixes_windows_reserved_device_names():
    assert xhs.safe_path_part("CON") == "_CON"
    assert xhs.safe_path_part("NUL.txt") == "_NUL.txt"


def test_supported_image_content_accepts_only_jpeg_and_png_magic_numbers():
    assert xhs.is_supported_image_content("image/jpeg", b"\xff\xd8\xff\xe0jpeg")
    assert xhs.is_supported_image_content("image/png", b"\x89PNG\r\n\x1a\npng")
    assert xhs.is_supported_image_content("", b"\xff\xd8\xff\xe0jpeg")
    assert xhs.is_supported_image_content("application/octet-stream", b"\x89PNG\r\n\x1a\npng")

    assert not xhs.is_supported_image_content("image/webp", b"RIFFxxxxWEBP")
    assert not xhs.is_supported_image_content("image/svg+xml", b"<svg></svg>")
    assert not xhs.is_supported_image_content("image/jpeg", b"not actually jpeg")


def test_http_downloader_accept_header_prefers_jpeg_and_png():
    downloader = xhs.HttpImageDownloader()
    try:
        assert downloader._client.headers["accept"] == "image/jpeg,image/png,*/*;q=0.8"
    finally:
        downloader.close()


def test_normalize_candidate_image_url_filters_noise_and_dedupes_variants():
    first = xhs.normalize_candidate_image_url(
        "https://sns-img-qc.xhscdn.com/photo.jpg?imageView2/2/w/1080/format/jpg#crop"
    )
    second = xhs.normalize_candidate_image_url(
        "https://sns-img-qc.xhscdn.com/photo.jpg?xsec_token=abc&imageView2/2/w/720"
    )

    assert first == "https://sns-img-qc.xhscdn.com/photo.jpg?imageView2/2/w/1080/format/jpg"
    assert xhs.image_url_dedupe_key(first) == xhs.image_url_dedupe_key(second)
    assert xhs.normalize_candidate_image_url("data:image/png;base64,abc") is None
    assert xhs.normalize_candidate_image_url("https://example.com/avatar.png") is None
    assert xhs.normalize_candidate_image_url("https://example.com/icon.svg") is None
    assert xhs.normalize_candidate_image_url("") is None


def test_existing_image_numbers_count_only_target_jpgs(tmp_path: Path):
    for name in [
        "01.jpg",
        "07.jpeg",
        "1.jpg",
        "2.jpg",
        "001.jpg",
        "08.jpg",
        "00.jpeg",
        "01.png",
        "note.jpg",
    ]:
        (tmp_path / name).write_bytes(b"image")

    assert xhs.existing_image_numbers(tmp_path) == {1, 7}


def test_next_image_number_returns_none_when_full(tmp_path: Path):
    for number in range(1, xhs.TARGET_IMAGES_PER_MODEL + 1):
        (tmp_path / f"{number:02d}.jpg").write_bytes(b"image")

    assert xhs.next_image_number(tmp_path) is None


def test_save_image_bytes_starts_with_zero_padded_one(tmp_path: Path):
    saved_path = xhs.save_image_bytes(tmp_path, b"new image")

    assert saved_path == tmp_path / "01.jpg"
    assert saved_path.read_bytes() == b"new image"


def test_save_image_bytes_uses_next_available_zero_padded_number(tmp_path: Path):
    (tmp_path / "01.jpg").write_bytes(b"old")
    (tmp_path / "03.jpeg").write_bytes(b"old")

    saved_path = xhs.save_image_bytes(tmp_path, b"new image")

    assert saved_path == tmp_path / "02.jpg"
    assert saved_path.read_bytes() == b"new image"


def test_save_image_bytes_returns_none_when_full_without_adding_files(tmp_path: Path):
    for number in range(1, xhs.TARGET_IMAGES_PER_MODEL + 1):
        (tmp_path / f"{number:02d}.jpg").write_bytes(b"image")
    before = sorted(path.name for path in tmp_path.iterdir())

    saved_path = xhs.save_image_bytes(tmp_path, b"new image")

    assert saved_path is None
    assert sorted(path.name for path in tmp_path.iterdir()) == before


class FakeBrowser:
    def __init__(self, results_by_query: dict[str, list[str] | Exception]):
        self.results_by_query = results_by_query
        self.queries: list[tuple[str, int]] = []

    def image_urls_for_query(self, query: str, limit: int) -> list[str]:
        self.queries.append((query, limit))
        result = self.results_by_query[query]
        if isinstance(result, Exception):
            raise result
        return result


class FakeDownloader:
    def __init__(self, payloads_by_url: dict[str, bytes | Exception | None]):
        self.payloads_by_url = payloads_by_url
        self.urls: list[str] = []

    def download(self, url: str) -> bytes | None:
        self.urls.append(url)
        payload = self.payloads_by_url[url]
        if isinstance(payload, Exception):
            raise payload
        return payload


class NoSleepRateLimiter:
    def __init__(self):
        self.searches = 0
        self.note_opens = 0
        self.image_saves = 0
        self.models: list[int] = []

    def wait(self) -> None:
        self.after_image_save()

    def after_search(self) -> None:
        self.searches += 1

    def after_note_open(self) -> None:
        self.note_opens += 1

    def after_image_save(self) -> None:
        self.image_saves += 1

    def after_model(self, processed_count: int) -> None:
        self.models.append(processed_count)


class FakeNoteBrowser(xhs.XhsBrowser):
    def __init__(self):
        self.opened: list[str] = []
        self._rate_limiter = NoSleepRateLimiter()

    def ensure_login(self) -> None:
        pass

    def _submit_search(self, query: str) -> None:
        pass

    def _collect_note_links(self, limit: int) -> list[str]:
        return ["https://www.xiaohongshu.com/explore/broken", "https://www.xiaohongshu.com/explore/ok"]

    def _image_urls_from_note(self, note_url: str) -> list[str]:
        self.opened.append(note_url)
        if note_url.endswith("/broken"):
            raise RuntimeError("note parse failed")
        return ["https://sns-img-qc.xhscdn.com/ok.jpg"]


def test_rate_limiter_waits_for_expected_runtime_events():
    sleeps: list[float] = []

    def choose_lower_bound(low: float, high: float) -> float:
        return low

    limiter = xhs.RateLimiter(sleep=sleeps.append, random_uniform=choose_lower_bound)

    limiter.after_search()
    limiter.after_note_open()
    limiter.after_image_save()
    limiter.after_model(9)
    limiter.after_model(10)

    assert sleeps == [3.0, 2.0, 0.5, 60.0]


def test_browser_skips_failed_note_and_keeps_collecting_images(capsys):
    browser = FakeNoteBrowser()

    urls = browser.image_urls_for_query("query", 1)

    assert urls == ["https://sns-img-qc.xhscdn.com/ok.jpg"]
    assert browser.opened == [
        "https://www.xiaohongshu.com/explore/broken",
        "https://www.xiaohongshu.com/explore/ok",
    ]
    assert browser._rate_limiter.note_opens == 2
    output = capsys.readouterr().out
    assert "跳过笔记：https://www.xiaohongshu.com/explore/broken，原因：note parse failed" in output


def make_task(model_name: str, search_query: str, output_dir: Path):
    return xhs.ModelTask(
        brand="Canon",
        brand_dir="Canon",
        model_name=model_name,
        canonical_name=model_name,
        common_name=model_name,
        search_query=search_query,
        output_dir=output_dir,
    )


def test_process_tasks_skips_full_folder_and_saves_until_seven(tmp_path: Path, capsys):
    full_dir = tmp_path / "full"
    partial_dir = tmp_path / "partial"
    full_dir.mkdir()
    partial_dir.mkdir()
    for number in range(1, xhs.TARGET_IMAGES_PER_MODEL + 1):
        (full_dir / f"{number:02d}.jpg").write_bytes(b"existing")
    (partial_dir / "01.jpg").write_bytes(b"existing")
    (partial_dir / "03.jpeg").write_bytes(b"existing")

    tasks = [
        make_task("full-model", "full query", full_dir),
        make_task("partial-model", "partial query", partial_dir),
    ]
    browser = FakeBrowser(
        {
            "partial query": [
                "none",
                "error",
                "image-1",
                "image-2",
                "image-3",
                "image-4",
                "image-5",
                "unused",
            ]
        }
    )
    downloader = FakeDownloader(
        {
            "none": None,
            "error": RuntimeError("network failed"),
            "image-1": b"image 1",
            "image-2": b"image 2",
            "image-3": b"image 3",
            "image-4": b"image 4",
            "image-5": b"image 5",
            "unused": b"unused",
        }
    )
    rate_limiter = NoSleepRateLimiter()

    summary = xhs.process_tasks(tasks, browser, downloader, rate_limiter)

    assert summary.total == 2
    assert summary.skipped_full == 1
    assert summary.searched == 1
    assert summary.saved == 5
    assert summary.incomplete == 0
    assert summary.failed == 0
    assert len(summary.failed_downloads) == 2
    assert browser.queries == [("partial query", xhs.TARGET_IMAGES_PER_MODEL)]
    assert downloader.urls == ["none", "error", "image-1", "image-2", "image-3", "image-4", "image-5"]
    assert rate_limiter.searches == 1
    assert rate_limiter.image_saves == 5
    assert rate_limiter.models == [1]
    assert sorted(path.name for path in partial_dir.iterdir()) == [
        "01.jpg",
        "02.jpg",
        "03.jpeg",
        "04.jpg",
        "05.jpg",
        "06.jpg",
        "07.jpg",
    ]
    output = capsys.readouterr().out
    assert "SKIP full-model" in output
    assert "FAIL partial-model none" in output
    assert "SAVE partial-model 07.jpg" in output


def test_run_summary_exposes_task_six_attribute_names():
    summary = xhs.RunSummary(total=1, skipped_full=2, saved=3, incomplete=4, failed=5)

    assert summary.total == 1
    assert summary.skipped_full == 2
    assert summary.saved == 3
    assert summary.incomplete == 4
    assert summary.failed == 5


def test_process_tasks_counts_empty_urls_as_incomplete(tmp_path: Path):
    task = make_task("empty-model", "empty query", tmp_path / "empty")
    browser = FakeBrowser({"empty query": []})
    downloader = FakeDownloader({})
    rate_limiter = NoSleepRateLimiter()

    summary = xhs.process_tasks([task], browser, downloader, rate_limiter)

    assert summary.total == 1
    assert summary.saved == 0
    assert summary.incomplete == 1
    assert summary.failed == 0
    assert downloader.urls == []
    assert rate_limiter.searches == 1
    assert rate_limiter.image_saves == 0
    assert rate_limiter.models == [1]


def test_process_tasks_prints_shortfall_without_marking_failed(tmp_path: Path, capsys):
    task = make_task("short-model", "short query", tmp_path / "short")
    browser = FakeBrowser({"short query": ["image-1", "image-2"]})
    downloader = FakeDownloader({"image-1": b"image 1", "image-2": b"image 2"})
    rate_limiter = NoSleepRateLimiter()

    summary = xhs.process_tasks([task], browser, downloader, rate_limiter)

    assert summary.incomplete == 1
    assert summary.failed == 0
    output = capsys.readouterr().out
    assert "short-model" in output
    assert "当前数量=2" in output
    assert "缺口=5" in output


def test_incomplete_summary_keeps_success_exit_code():
    assert xhs.exit_code_for_summary(xhs.RunSummary(incomplete=1, failed=0)) == 0
    assert xhs.exit_code_for_summary(xhs.RunSummary(incomplete=0, failed=1)) == 1


def test_process_tasks_counts_search_error_as_failed_and_continues(tmp_path: Path):
    first_dir = tmp_path / "first"
    second_dir = tmp_path / "second"
    tasks = [
        make_task("broken-model", "broken query", first_dir),
        make_task("next-model", "next query", second_dir),
    ]
    browser = FakeBrowser(
        {
            "broken query": RuntimeError("search failed"),
            "next query": [
                "image-1",
                "image-2",
                "image-3",
                "image-4",
                "image-5",
                "image-6",
                "image-7",
            ],
        }
    )
    downloader = FakeDownloader(
        {
            "image-1": b"image 1",
            "image-2": b"image 2",
            "image-3": b"image 3",
            "image-4": b"image 4",
            "image-5": b"image 5",
            "image-6": b"image 6",
            "image-7": b"image 7",
        }
    )
    rate_limiter = NoSleepRateLimiter()

    summary = xhs.process_tasks(tasks, browser, downloader, rate_limiter)

    assert summary.total == 2
    assert summary.failed == 1
    assert summary.incomplete == 1
    assert summary.saved == 7
    assert sorted(path.name for path in second_dir.iterdir()) == [
        "01.jpg",
        "02.jpg",
        "03.jpg",
        "04.jpg",
        "05.jpg",
        "06.jpg",
        "07.jpg",
    ]
